"""Phase 0 brief: extracting an old manual / outline / process document, and how the brief flows into
config, manual.yaml, rendering and the run report."""

from __future__ import annotations

import http.server
import io
import json
import subprocess
import sys
import threading
from pathlib import Path

import pytest
from docx import Document

from guidegen_engine import brief
from guidegen_engine.brief import extract
from guidegen_engine.config import load_config
from guidegen_engine.configinit import set_value
from guidegen_engine.errors import GuideGenError
from guidegen_engine.manual import ManualStore
from guidegen_engine.masking import register_secret
from guidegen_engine.render import render
from guidegen_engine.report import write_report

from conftest import write_config
from test_render import Links, _captured

EMAIL = r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"


def _old_manual_docx(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("Acme Quote Manual", style="Title")
    doc.add_paragraph("Getting started", style="Heading 1")
    doc.add_paragraph("Contact jane.real@corp.example for an account.")
    doc.add_paragraph("Creating quotes", style="Heading 1")
    doc.add_paragraph("Create a quote", style="Heading 2")
    doc.add_paragraph("Open the Quotes tab.", style="List Number")
    p = doc.add_paragraph(style="List Number")
    p.add_run("Click ")
    p.add_run("New quote").bold = True
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Term", "Meaning"
    t.cell(1, 0).text, t.cell(1, 1).text = "Quote", "An offer to a customer"
    doc.save(str(path))
    return path


def test_extract_docx_outline_lists_tables_and_redaction(ctx):
    src = _old_manual_docx(ctx.repo / "old-manual.docx")
    res = extract(ctx, "old-manual.docx", [EMAIL])
    assert res["ok"] and res["format"] == "docx"
    titles = [(h["level"], h["title"]) for h in res["outline"]]
    assert titles == [(1, "Acme Quote Manual"), (1, "Getting started"), (1, "Creating quotes"), (2, "Create a quote")]
    text = Path(res["file"]).read_text(encoding="utf-8")
    assert "1. Open the Quotes tab." in text and "**New quote**" in text
    assert "| Quote | An offer to a customer |" in text
    assert "jane.real@corp.example" not in text and "[redacted]" in text and res["redactions"] == 1
    assert Path(res["file"]).parent == ctx.work / "brief"
    assert src.exists()  # the source is only read


def test_extract_markdown_and_html(ctx):
    (ctx.repo / "outline.md").write_text("# Manual plan\n\n## Orders\n- Create an order\n\n```\n# not a heading\n```\n## Glossary\n", encoding="utf-8")
    res = extract(ctx, "outline.md")
    assert [h["title"] for h in res["outline"]] == ["Manual plan", "Orders", "Glossary"]

    (ctx.repo / "help.html").write_text(
        "<html><head><title>x</title><style>h1{}</style></head><body><h1>Help</h1><p>Use <b>Save</b>.</p>"
        "<script>ignore()</script><ol><li>One</li><li>Two</li></ol><h2>FAQ</h2></body></html>", encoding="utf-8")
    res = extract(ctx, "help.html")
    text = Path(res["file"]).read_text(encoding="utf-8")
    assert [h["title"] for h in res["outline"]] == ["Help", "FAQ"]
    assert "Use **Save**." in text and "1. One" in text and "ignore()" not in text


def test_extract_pdf_text_and_bookmarks(ctx):
    pypdf = pytest.importorskip("pypdf")
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (200, 100), "white").save(buf, format="PDF")
    w = pypdf.PdfWriter(clone_from=pypdf.PdfReader(io.BytesIO(buf.getvalue())))
    w.add_outline_item("Chapter 1: Orders", 0)
    out = ctx.repo / "old.pdf"
    with out.open("wb") as f:
        w.write(f)
    res = extract(ctx, str(out))  # absolute path works too
    assert res["format"] == "pdf"
    assert res["outline"][0]["title"] == "Chapter 1: Orders" and res["outline"][0]["page"] == 1
    assert "hint" in res  # image-only PDF: almost no text


def test_extract_url(ctx):
    class H(http.server.BaseHTTPRequestHandler):
        def do_GET(self):  # noqa: N802
            body = b"<h1>Online help</h1><p>Call support.</p>"
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def log_message(self, *a):
            pass

    srv = http.server.HTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        res = extract(ctx, f"http://127.0.0.1:{srv.server_address[1]}/help/start")
    finally:
        srv.shutdown()
    assert res["format"] == "url/html" and res["outline"][0]["title"] == "Online help"


def test_extract_errors_and_secret_masking(ctx, monkeypatch):
    with pytest.raises(GuideGenError, match="not found"):
        extract(ctx, "missing.docx")
    (ctx.repo / "a.doc").write_bytes(b"\xd0\xcf")
    with pytest.raises(GuideGenError, match=r"\.docx or PDF"):
        extract(ctx, "a.doc")
    (ctx.repo / "a.xyz").write_text("x", encoding="utf-8")
    with pytest.raises(GuideGenError, match="unsupported"):
        extract(ctx, "a.xyz")
    (ctx.repo / "broken.docx").write_bytes(b"not a zip")
    with pytest.raises(GuideGenError, match="docx"):
        extract(ctx, "broken.docx")
    monkeypatch.setattr(brief, "MAX_BYTES", 10)
    (ctx.repo / "big.md").write_text("x" * 50, encoding="utf-8")
    with pytest.raises(GuideGenError, match="larger"):
        extract(ctx, "big.md")
    monkeypatch.setattr(brief, "MAX_BYTES", 20 * 1024 * 1024)
    register_secret("hunter2-secret")
    (ctx.repo / "notes.txt").write_text("the password is hunter2-secret", encoding="utf-8")
    res = extract(ctx, "notes.txt")
    assert "hunter2-secret" not in Path(res["file"]).read_text(encoding="utf-8")


def test_brief_config(ctx, cfg):
    write_config(ctx, cfg)
    assert load_config(ctx.config_path).brief.asked is False
    set_value(ctx, "brief", json.dumps({"asked": True, "sources": [
        {"kind": "old-manual", "path": "docs/old-manual.docx"},
        {"kind": "reference-manual", "url": "https://help.example.com/guide"},
    ], "notes": "Audience: warehouse staff. Say 'item', never 'SKU'."}))
    b = load_config(ctx.config_path).brief
    assert b.asked and b.sources[1].url.startswith("https://") and "warehouse" in b.notes
    with pytest.raises(GuideGenError):
        set_value(ctx, "brief.sources", '[{"kind": "novel", "path": "x"}]')
    with pytest.raises(GuideGenError):
        set_value(ctx, "brief.sources", '[{"kind": "outline", "url": "file:///etc/passwd"}]')


def test_custom_chapters_merge_keeps_developer_edits(tmp_path):
    store = ManualStore.load(tmp_path / "manual.yaml")
    store.merge({"chapters": [{"id": "about"}, {"id": "getting-started"}, {"id": "tasks"}, {"id": "glossary", "title": "Glossary", "body": "- **Quote**: an offer"}]})
    store.save()
    # the developer overrides the glossary and hides the about chapter
    store = ManualStore.load(tmp_path / "manual.yaml")
    store.data["chapters"][3]["override"] = "My glossary."
    store.data["chapters"][0]["exclude"] = True
    store.save()
    # regeneration re-sends chapters in a new order: order follows the patch, developer fields survive
    store = ManualStore.load(tmp_path / "manual.yaml")
    store.merge({"chapters": [{"id": "getting-started"}, {"id": "glossary", "body": "- **Quote**: new text"}, {"id": "about", "body": "Hi"}, {"id": "tasks"}]})
    store.save()
    m = ManualStore.load(tmp_path / "manual.yaml").model()
    assert [c.id for c in m.chapters] == ["getting-started", "glossary", "about", "tasks"]
    gl = m.chapters[1]
    assert gl.text == "My glossary." and gl.body == "- **Quote**: new text" and m.chapters[2].exclude
    store.merge({"chapters": [{"id": "tasks"}, {"id": "tasks"}]})
    with pytest.raises(GuideGenError, match="duplicate"):
        store.model()
    bad = ManualStore.load(tmp_path / "x.yaml")
    bad.merge({"chapters": [{"id": "Bad Id"}]})
    with pytest.raises(GuideGenError):
        bad.model()


def _briefed(ctx, fake):
    cfg, _ = _captured(ctx, fake, brief={"asked": True, "sources": [{"kind": "old-manual", "path": "old-manual.docx"}]})
    store = ManualStore.load(ctx.manual_path)
    store.merge({
        "meta": {"audience": "sales staff"},
        "chapters": [{"id": "about", "title": "About this manual", "body": "Written for **sales staff**."},
                     {"id": "getting-started"}, {"id": "tasks"}, {"id": "reference"},
                     {"id": "glossary", "title": "Glossary", "body": "- **Quote**: an offer to a customer"},
                     {"id": "troubleshooting"}, {"id": "empty"}],
        "tasks": [{"id": "create-order", "group": "Orders", "evidence": [{"type": "brief", "ref": "old-manual.docx §Create a quote"}]}],
        "briefGaps": [{"item": "Print a quote", "source": "old-manual.docx §Printing", "reason": "no print feature in the app"}],
    })
    store.save()
    return cfg, ManualStore.load(ctx.manual_path).model()


def test_render_custom_chapters_groups_and_audience(ctx, fake):
    cfg, manual = _briefed(ctx, fake)
    render(ctx, cfg, manual, ["docx", "html"])

    html = (ctx.out / "manual-html" / "index.html").read_text(encoding="utf-8")
    p = Links()
    p.feed(html)
    assert not [h for h in p.hrefs if h not in p.ids]
    assert 'id="chapter-glossary"' in html and "an offer to a customer" in html
    assert html.index("About this manual") < html.index("Getting Started") < html.index("Glossary")
    assert '<h3 class="group">Orders</h3>' in html and "For sales staff" in html
    assert 'id="chapter-empty"' not in html  # empty custom chapters are skipped
    images = {f.name for f in (ctx.out / "manual-html" / "images").iterdir()}
    assert not [n for n in images if n.startswith(".")]

    doc = Document(str(ctx.out / "manual.docx"))
    heads = [(pp.style.name, pp.text) for pp in doc.paragraphs if pp.style.name.startswith("Heading")]
    h1 = [t for s, t in heads if s == "Heading 1"]
    assert h1.index("About this manual") < h1.index("Getting Started") < h1.index("Glossary") < h1.index("Troubleshooting")
    assert ("Heading 2", "Orders") in heads and ("Heading 3", "Create an order") in heads
    assert "For sales staff" in "\n".join(pp.text for pp in doc.paragraphs)


def test_report_brief_coverage(ctx, fake):
    cfg, manual = _briefed(ctx, fake)
    res = write_report(ctx, cfg, manual)
    text = (ctx.out / "run-report.md").read_text(encoding="utf-8")
    assert "## Your brief" in text and "old-manual.docx" in text and "Print a quote" in text
    assert res["brief"] == {"sources": 1, "tasksFromBrief": 1, "gaps": 1}
    assert any("brief" in i for i in res["topIssues"]) or len(res["topIssues"]) == 3


def test_cli_brief_extract_without_config(ctx):
    (ctx.repo / ".git").mkdir()
    (ctx.repo / "plan.md").write_text("# Plan\n## Orders\nmail me: a.b@c.example\n", encoding="utf-8")
    r = subprocess.run([sys.executable, "-m", "guidegen_engine", "brief", "extract", "plan.md", "--json", "--repo", str(ctx.repo)],
                       capture_output=True, text=True, cwd=ctx.repo, timeout=120)
    assert r.returncode == 0, r.stderr
    out = json.loads(r.stdout)
    assert out["ok"] and [h["title"] for h in out["outline"]] == ["Plan", "Orders"] and out["redactions"] == 1
