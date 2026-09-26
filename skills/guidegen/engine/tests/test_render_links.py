"""Cross-link, image-name and Word layout checks for the renderers (review findings)."""

from __future__ import annotations

import json
import re
import zipfile

from docx import Document
from PIL import Image

from guidegen_engine.manual.models import Manual
from guidegen_engine.render import render
from guidegen_engine.render.docx import _rich
from guidegen_engine.render.docx.figure import MAX_HEIGHT_IN, MAX_WIDTH_IN, fit_inches
from guidegen_engine.render.docx.ox import bookmark_name
from guidegen_engine.report import write_report

from conftest import base_config
from test_render import Links


def _png(path, size=(400, 300), color="white"):
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, color).save(path)


def _manual(ctx, **over) -> Manual:
    # screen "a-b" -> screens/a-b.png and task "a" step "b" -> screens/a/b.png used to flatten to one name
    _png(ctx.out / "screens" / "a-b.png", color="red")
    _png(ctx.out / "screens" / "a" / "b.png", color="blue")
    _png(ctx.out / "screens" / "order_list.png", color="green")
    data = {
        "screens": [
            {"id": "a-b", "title": "Dashboard", "status": "captured", "image": "screens/a-b.png"},
            {"id": "order_list", "title": "Orders", "status": "captured", "image": "screens/order_list.png"},
        ],
        "tasks": [
            {"id": "a", "title": "Open orders", "rank": 1, "status": "captured", "steps": [
                {"id": "b", "screen": "order_list", "status": "captured", "image": "screens/a/b.png",
                 "text": "Click **Orders**."},
            ]},
            {"id": "a-b-x", "title": "Other", "rank": 2, "status": "captured", "steps": [
                {"id": "s1", "screen": "a-b", "status": "captured", "text": "Look."},
            ]},
        ],
    }
    data.update(over)
    return Manual.model_validate(data)


def _docx_links_resolve(path) -> None:
    xml = zipfile.ZipFile(path).read("word/document.xml").decode()
    bookmarks = set(re.findall(r'w:bookmarkStart[^>]*w:name="([^"]+)"', xml))
    anchors = re.findall(r'w:hyperlink[^>]*w:anchor="([^"]+)"', xml)
    assert anchors, "expected internal links"
    missing = [a for a in anchors if a not in bookmarks]
    assert not missing, missing


def test_underscore_ids_link_and_images_do_not_collide(ctx):
    manual = _manual(ctx)
    render(ctx, base_config(), manual, ["html", "docx"])
    html = (ctx.out / "manual-html" / "index.html").read_text(encoding="utf-8")
    p = Links()
    p.feed(html)
    missing = [h for h in p.hrefs if h not in p.ids]
    assert not missing, missing
    assert "#screen-order_list" in html
    imgs = {s for s in p.imgs if s}
    assert "images/a-b.png" in imgs and "images/a/b.png" in imgs
    red = Image.open(ctx.out / "manual-html" / "images" / "a-b.png").getpixel((5, 5))
    blue = Image.open(ctx.out / "manual-html" / "images" / "a" / "b.png").getpixel((5, 5))
    assert red != blue
    # step anchor "step-a-b" must not clash with the task anchor of a task called "a-b"
    assert len(re.findall(r'id="task-a-b"', html)) <= 1
    _docx_links_resolve(ctx.out / "manual.docx")


def test_no_links_into_excluded_chapter(ctx):
    manual = _manual(ctx, chapters=[{"id": "getting-started"}, {"id": "tasks"}, {"id": "reference", "exclude": True}])
    render(ctx, base_config(), manual, ["html", "docx"])
    html = (ctx.out / "manual-html" / "index.html").read_text(encoding="utf-8")
    p = Links()
    p.feed(html)
    assert not [h for h in p.hrefs if h not in p.ids]
    assert 'href="#screen-' not in html
    xml = zipfile.ZipFile(ctx.out / "manual.docx").read("word/document.xml").decode()
    assert "b_screen_" not in xml


def test_bookmark_names_unique_and_valid():
    names = [bookmark_name(a) for a in ("task-a-b", "task-a_b", "screen-" + "x" * 60, "screen-" + "x" * 59 + "y")]
    assert len(set(names)) == len(names)
    assert all(len(n) <= 40 and re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", n) for n in names)


def test_figures_fit_the_page():
    assert fit_inches(1280, 4000) * 4000 / 1280 <= MAX_HEIGHT_IN + 1e-6
    assert fit_inches(1440, 900) == MAX_WIDTH_IN
    assert fit_inches(200, 100) == 200 / 96


def test_numbered_lists_restart():
    doc = Document()
    _rich(doc, "1. one\n2. two")
    _rich(doc, "Some text.")
    _rich(doc, "1. again\n2. more")
    ids = [p._p.pPr.numPr.numId.val for p in doc.paragraphs if p.style.name == "List Number"]
    assert len(ids) == 4 and ids[0] == ids[1] and ids[2] == ids[3] and ids[0] != ids[2]


def test_markdown_in_word_tables(ctx):
    manual = _manual(ctx, troubleshooting=[{"message": "Save failed", "cause": "The **name** is empty.", "fix": "- Enter a name\n- Click **Save**"}])
    render(ctx, base_config(), manual, ["docx"])
    doc = Document(str(ctx.out / "manual.docx"))
    cells = [c for t in doc.tables for r in t.rows for c in r.cells]
    text = "\n".join(c.text for c in cells)
    assert "**" not in text
    assert any(r.bold and r.text == "Save" for c in cells for p in c.paragraphs for r in p.runs)


def test_report_finds_single_file_html(ctx):
    manual = _manual(ctx)
    cfg = base_config()
    render(ctx, cfg, manual, ["html"], single_file=True)
    assert not (ctx.out / "manual-html").exists()
    rep = write_report(ctx, cfg, manual)
    assert rep["outputs"]["html"] and rep["outputs"]["html"].endswith("manual.html")


def test_search_index_only_lists_rendered_sections(ctx):
    manual = _manual(ctx, chapters=[{"id": "getting-started"}, {"id": "tasks", "exclude": True}, {"id": "reference"}])
    render(ctx, base_config(), manual, ["html"])
    html = (ctx.out / "manual-html" / "index.html").read_text(encoding="utf-8")
    index = json.loads(re.search(r'<script[^>]*id="search-index"[^>]*>(.*?)</script>', html, re.S).group(1))
    p = Links()
    p.feed(html)
    targets = [e["u"][1:] for e in index]
    assert targets and all(t in p.ids for t in targets), targets
