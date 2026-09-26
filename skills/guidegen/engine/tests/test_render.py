from __future__ import annotations

import re
import zipfile
from html.parser import HTMLParser

from docx import Document

from guidegen_engine.capture import run_capture
from guidegen_engine.manual import ManualStore
from guidegen_engine.render import render
from guidegen_engine.report import write_report
from guidegen_engine.runtime import Session

from conftest import base_config
from test_capture import MANUAL


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.ids, self.hrefs, self.imgs, self.alts_missing = set(), [], [], 0

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if "id" in a:
            self.ids.add(a["id"])
        if tag == "a" and a.get("href", "").startswith("#"):
            self.hrefs.append(a["href"][1:])
        if tag == "img":
            self.imgs.append(a.get("src"))
            if "alt" not in a:
                self.alts_missing += 1


def _captured(ctx, fake, **over):
    store = ManualStore.load(ctx.manual_path)
    store.merge(MANUAL)
    store.merge({
        "getting_started": {"launch": "Open **Acme Quote** from the Start menu.", "tour": "The home screen shows your work."},
        "troubleshooting": [{"message": "Quantity must be greater than 0", "source": "src/OrderValidator.cs",
                             "cause": "The quantity was 0.", "fix": "Enter **1** or more."}],
    })
    store.merge({"screens": [{"id": "orders-list", "title": "Orders", "text": {"summary": "All orders.", "override": "Custom orders text."},
                              "elements": [{"name": "New Order", "type": "button", "description": "Opens the **New Order** dialog."}]}]})
    store.save()
    cfg = base_config(**over)
    run_capture(ctx, Session(cfg, fake), ManualStore.load(ctx.manual_path), progress=lambda m: None)
    return cfg, ManualStore.load(ctx.manual_path).model()


def test_render_html_links_and_offline(ctx, fake):
    cfg, manual = _captured(ctx, fake)
    res = render(ctx, cfg, manual, ["html"])
    html = (ctx.out / "manual-html" / "index.html").read_text(encoding="utf-8")
    p = Links()
    p.feed(html)
    missing = [h for h in p.hrefs if h not in p.ids]
    assert not missing, missing
    assert p.alts_missing == 0
    for src in p.imgs:
        if src:  # the zoom dialog img is filled at runtime
            assert (ctx.out / "manual-html" / src).exists(), src
    assert "http://" not in html and "https://" not in html  # works from file:// with no network
    assert "Custom orders text." in html and "All orders." not in html  # override wins
    assert "<strong>New Order</strong>" in html
    assert res["tasks"] == 3 and res["screens"] == 4


def test_single_file_html(ctx, fake):
    cfg, manual = _captured(ctx, fake)
    render(ctx, cfg, manual, ["html"], single_file=True)
    html = (ctx.out / "manual.html").read_text(encoding="utf-8")
    assert "data:image/png;base64," in html and 'src="images/' not in html


def test_render_docx(ctx, fake):
    cfg, manual = _captured(ctx, fake)
    render(ctx, cfg, manual, ["docx"])
    path = ctx.out / "manual.docx"
    doc = Document(str(path))
    headings = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert ("Heading 1", "How-To Tasks") in headings
    assert ("Heading 2", "Create an order") in headings
    assert ("Heading 1", "Screen Reference") in headings
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "Custom orders text." in body
    xml = zipfile.ZipFile(path).read("word/document.xml").decode()
    assert "TOC \\o" in xml and 'descr="' in xml and "w:bookmarkStart" in xml
    footer = zipfile.ZipFile(path).read("word/footer1.xml").decode() if "word/footer1.xml" in zipfile.ZipFile(path).namelist() else ""
    assert "NUMPAGES" in (footer + "".join(zipfile.ZipFile(path).read(n).decode() for n in zipfile.ZipFile(path).namelist() if n.startswith("word/footer")))


def test_low_contrast_brand_darkened_and_reported(ctx, fake):
    cfg, manual = _captured(ctx, fake, brand={"productName": "Pale", "primaryColor": "#9FD3FF"})
    res = render(ctx, cfg, manual, ["html"])
    assert res["notes"] and "darkened" in res["notes"][0]
    html = (ctx.out / "manual-html" / "index.html").read_text(encoding="utf-8")
    m = re.search(r"--primary: (#[0-9A-F]{6})", html)
    assert m and m.group(1) != "#9FD3FF"
    rep = write_report(ctx, cfg, manual)
    text = (ctx.out / "run-report.md").read_text(encoding="utf-8")
    assert "darkened" in text
    assert rep["screens"]["unreachable"] == 1 and "ghost" in text
    assert "Refused destructive actions" in text and rep["refused"] == 1
